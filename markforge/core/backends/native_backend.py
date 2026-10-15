"""Запасной backend на чистом Python (без pandoc).

Нужен для окружений, где pandoc недоступен (например, будущий сервер/мобилка).
Честные границы качества (см. таблицу в ARCHITECTURE.md):

* Хорошо: md↔html, docx→md, epub→md, md→epub, любое →pdf.
* Базово: md→docx, html→docx (заголовки/абзацы/жирный-курсив/списки/таблицы/картинки).
* docx↔epub оставлены pandoc'у (native не регистрирует эти пары).
"""

from __future__ import annotations

import itertools
import posixpath
import re
import uuid
from pathlib import Path
from typing import Dict, Iterator, List, Optional, Set, Tuple

from bs4 import BeautifulSoup

from markforge.core.backends.base import BackendOutput
from markforge.core.backends.pdf import render_html_to_pdf
from markforge.core.errors import ConversionFailedError, CorruptInputError
from markforge.core.formats import Format
from markforge.core.options import ConvertOptions

_IMG_EXT = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/jpg": ".jpg",
    "image/gif": ".gif",
    "image/svg+xml": ".svg",
    "image/webp": ".webp",
    "image/bmp": ".bmp",
    "image/tiff": ".tif",
}


def _read_text(path: Path) -> str:
    # utf-8-sig снимает BOM (частый на Windows), не мешая обычному UTF-8.
    return Path(path).read_text(encoding="utf-8-sig", errors="replace")


class NativeBackend:
    name = "native"

    def is_available(self) -> bool:
        return True

    def supports(self, src: Format, dst: Format) -> bool:
        return (src, dst) in _DISPATCH or (dst == Format.PDF and src in _PDF_SOURCES)

    def convert(
        self,
        input_path: Path,
        work_dir: Path,
        out_name: str,
        src: Format,
        dst: Format,
        options: ConvertOptions,
    ) -> BackendOutput:
        input_path = Path(input_path).resolve()
        out_path = work_dir / out_name

        if dst == Format.PDF:
            return _to_pdf(input_path, out_path, src, options)

        handler = _DISPATCH.get((src, dst))
        if handler is None:
            raise ConversionFailedError(
                f"native backend не умеет {src.title} → {dst.title}."
            )
        return handler(input_path, work_dir, out_path, options)


# --------------------------------------------------------------------------- #
# Markdown <-> HTML
# --------------------------------------------------------------------------- #
def _split_front_matter(text: str) -> Tuple[Dict[str, str], str]:
    """Отделить YAML-front matter (``---\\n...\\n---``) от тела Markdown."""
    meta: Dict[str, str] = {}
    if text.lstrip().startswith("---"):
        text = text.lstrip("﻿")
        m = re.match(r"^---[ \t]*\n(.*?)\n(?:---|\.\.\.)[ \t]*\n", text, re.DOTALL)
        if m:
            for line in m.group(1).splitlines():
                if ":" in line and not line.lstrip().startswith("#"):
                    key, _, val = line.partition(":")
                    meta[key.strip().lower()] = val.strip().strip("\"'")
            text = text[m.end():]
    return meta, text


def _md_to_html_str(md_text: str) -> str:
    import markdown

    _, body = _split_front_matter(md_text)
    return markdown.markdown(
        body,
        extensions=["extra", "sane_lists", "toc", "admonition"],
        output_format="html5",
    )


def _html_body(html: str) -> BeautifulSoup:
    soup = BeautifulSoup(html, "html.parser")
    return soup.body or soup


def _md_to_html(inp: Path, work_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    html_body = _md_to_html_str(_read_text(inp))
    if o.standalone:
        title = _first_heading(html_body) or inp.stem
        doc = (
            "<!DOCTYPE html><html><head><meta charset='utf-8'>"
            f"<title>{title}</title></head><body>{html_body}</body></html>"
        )
    else:
        doc = html_body
    out_path.write_text(doc, encoding="utf-8")
    return BackendOutput()


def _html_to_md(inp: Path, work_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    from markdownify import markdownify as to_md

    body = _html_body(_read_text(inp))
    md_text = to_md(str(body), heading_style="ATX", bullets="-")
    out_path.write_text(_clean_md(md_text), encoding="utf-8")
    return BackendOutput()


# --------------------------------------------------------------------------- #
# DOCX -> Markdown  (mammoth -> html -> markdownify, картинки в media/)
# --------------------------------------------------------------------------- #
def _docx_to_md(inp: Path, work_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    import mammoth
    from markdownify import markdownify as to_md

    media_dir = work_dir / o.media_dir_name
    counter = itertools.count(1)
    warnings: List[str] = []

    def handler(image):
        ext = _IMG_EXT.get(image.content_type, ".png")
        name = f"image{next(counter)}{ext}"
        media_dir.mkdir(parents=True, exist_ok=True)
        with image.open() as stream:
            (media_dir / name).write_bytes(stream.read())
        return {"src": f"{o.media_dir_name}/{name}"}

    try:
        with open(inp, "rb") as fh:
            result = mammoth.convert_to_html(
                fh, convert_image=mammoth.images.img_element(handler)
            )
    except Exception as exc:  # noqa: BLE001
        raise CorruptInputError(f"Не удалось прочитать DOCX: {exc}") from exc

    md_text = to_md(result.value, heading_style="ATX", bullets="-")
    out_path.write_text(_clean_md(md_text), encoding="utf-8")
    warnings += [m.message for m in result.messages if getattr(m, "type", "") == "warning"]
    return BackendOutput(warnings=warnings)


# --------------------------------------------------------------------------- #
# EPUB -> Markdown  (ebooklib -> html per spine doc -> markdownify)
# --------------------------------------------------------------------------- #
def _load_epub(inp: Path):
    import ebooklib
    from ebooklib import epub

    try:
        return ebooklib, epub, epub.read_epub(str(inp))
    except Exception as exc:  # noqa: BLE001
        raise CorruptInputError(f"Не удалось прочитать EPUB: {exc}") from exc


def _norm_epub_path(name: str) -> str:
    return posixpath.normpath(name.replace("\\", "/")).lstrip("/")


def _unique_media_name(base: str, used: Set[str]) -> str:
    """Уникальное имя файла в media/ (разрешает коллизии basename из разных папок)."""
    if base not in used:
        return base
    stem, dot, ext = base.rpartition(".")
    if not dot:
        stem, ext = base, ""
    else:
        ext = "." + ext
    i = 1
    while f"{stem}_{i}{ext}" in used:
        i += 1
    return f"{stem}_{i}{ext}"


def _resolve_doc_img(doc_name: str, src: str) -> str:
    """Полный нормализованный путь картинки внутри EPUB (src относителен документу)."""
    doc_dir = posixpath.dirname(_norm_epub_path(doc_name))
    return posixpath.normpath(posixpath.join(doc_dir, src.replace("\\", "/"))).lstrip("/")


def _epub_images_to_media(ebooklib, book, media_dir: Path) -> Dict[str, str]:
    """Выгрузить картинки EPUB в media/, вернуть карту полный_путь -> имя_в_media."""
    mapping: Dict[str, str] = {}
    used: Set[str] = set()
    for item in book.get_items_of_type(ebooklib.ITEM_IMAGE):
        full = _norm_epub_path(item.get_name())
        name = _unique_media_name(posixpath.basename(full) or "image", used)
        used.add(name)
        media_dir.mkdir(parents=True, exist_ok=True)
        (media_dir / name).write_bytes(item.get_content())
        mapping[full] = name
    return mapping


def _lookup_by_path(mapping: Dict[str, str], full: str):
    """Найти значение по полному пути, с откатом на совпадение basename."""
    if full in mapping:
        return mapping[full]
    base = posixpath.basename(full)
    for key, value in mapping.items():
        if posixpath.basename(key) == base:
            return value
    return None


def _epub_to_md(inp: Path, work_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    from markdownify import markdownify as to_md

    ebooklib, epub, book = _load_epub(inp)
    media_dir = work_dir / o.media_dir_name
    mapping = _epub_images_to_media(ebooklib, book, media_dir)

    parts: List[str] = []
    for item in _spine_documents(ebooklib, book):
        html = item.get_content().decode("utf-8", "replace")
        body = _html_body(html)
        _rewrite_img_src(body, item.get_name(), mapping, o.media_dir_name)
        parts.append(to_md(str(body), heading_style="ATX", bullets="-"))

    md_text = _clean_md("\n\n".join(p for p in parts if p.strip()))
    out_path.write_text(md_text, encoding="utf-8")
    return BackendOutput()


def _spine_documents(ebooklib, book) -> Iterator:
    seen = set()
    for entry in book.spine:
        idref = entry[0] if isinstance(entry, (tuple, list)) else entry
        item = book.get_item_with_id(idref)
        if item is not None and item.get_type() == ebooklib.ITEM_DOCUMENT:
            seen.add(item.get_id())
            yield item
    # На случай кривого spine — добавим оставшиеся документы.
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        if item.get_id() not in seen:
            yield item


def _rewrite_img_src(body: BeautifulSoup, doc_name: str, mapping: Dict[str, str],
                     media_dir_name: str) -> None:
    for img in body.find_all("img"):
        src = img.get("src")
        if not src or src.startswith(("http://", "https://", "data:")):
            continue
        full = _resolve_doc_img(doc_name, src)
        name = _lookup_by_path(mapping, full)
        if name:
            img["src"] = f"{media_dir_name}/{name}"


# --------------------------------------------------------------------------- #
# Markdown -> EPUB  (md -> html -> ebooklib, картинки встраиваются)
# --------------------------------------------------------------------------- #
def _md_to_epub(inp: Path, work_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    from ebooklib import epub

    raw = _read_text(inp)
    meta, _ = _split_front_matter(raw)
    html_body = _md_to_html_str(raw)
    body = BeautifulSoup(html_body, "html.parser")
    m = o.epub
    title = m.title or meta.get("title") or _first_heading(html_body) or inp.stem
    lang = m.language or meta.get("lang") or meta.get("language") or "ru"
    authors = list(m.authors) if m.authors else (
        [meta["author"]] if meta.get("author") else []
    )

    book = epub.EpubBook()
    book.set_identifier(m.identifier or f"markforge-{uuid.uuid4()}")
    book.set_title(title)
    book.set_language(lang)
    for author in authors:
        book.add_author(author)

    # Встраиваем локальные картинки.
    base_dir = inp.parent
    counter = itertools.count(1)
    for img in body.find_all("img"):
        src = img.get("src")
        if not src or src.startswith(("http://", "https://", "data:")):
            continue
        img_path = (base_dir / src).resolve()
        if not img_path.exists():
            continue
        internal = f"images/img{next(counter)}{img_path.suffix.lower()}"
        item = epub.EpubItem(
            uid=internal,
            file_name=internal,
            media_type=_guess_media_type(img_path.suffix),
            content=img_path.read_bytes(),
        )
        book.add_item(item)
        img["src"] = internal

    if m.cover_image:
        cover = Path(m.cover_image)
        if cover.exists():
            book.set_cover("cover" + cover.suffix.lower(), cover.read_bytes())

    css = None
    if o.css and Path(o.css).exists():
        css = epub.EpubItem(
            uid="style", file_name="style/main.css",
            media_type="text/css", content=Path(o.css).read_bytes(),
        )
        book.add_item(css)

    # Проставляем уникальные id заголовкам, чтобы оглавление реально навигировало.
    toc_links = _assign_heading_ids(epub, body, "content.xhtml")

    chapter = epub.EpubHtml(title=title, file_name="content.xhtml", lang=lang)
    chapter.content = f"<html><body>{body}</body></html>"
    if css:
        chapter.add_item(css)
    book.add_item(chapter)

    book.toc = toc_links or [chapter]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]

    epub.write_epub(str(out_path), book)
    return BackendOutput()


def _assign_heading_ids(epub, body: BeautifulSoup, file_name: str) -> list:
    """Присвоить заголовкам уникальные id и вернуть ссылки TOC с якорями."""
    links = []
    used: Set[str] = set()
    n = 0
    for h in body.find_all(["h1", "h2"]):
        text = h.get_text(strip=True)
        if not text:
            continue
        n += 1
        hid = h.get("id") or f"h{n}"
        while hid in used:
            n += 1
            hid = f"h{n}"
        used.add(hid)
        h["id"] = hid
        links.append(epub.Link(f"{file_name}#{hid}", text, hid))
    return links


# --------------------------------------------------------------------------- #
# Markdown/HTML -> DOCX  (базовый построитель через python-docx)
# --------------------------------------------------------------------------- #
def _md_to_docx(inp: Path, work_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    html = _md_to_html_str(_read_text(inp))
    return _html_string_to_docx(html, inp.parent, out_path, o)


def _html_to_docx(inp: Path, work_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    body = _html_body(_read_text(inp))
    return _html_string_to_docx(str(body), inp.parent, out_path, o)


def _html_string_to_docx(html: str, base_dir: Path, out_path: Path, o: ConvertOptions) -> BackendOutput:
    from docx import Document
    from docx.shared import Pt

    doc = Document(str(o.reference_doc)) if o.reference_doc and Path(o.reference_doc).exists() else Document()
    body = _html_body(html)
    warnings: List[str] = [
        "Native backend: базовое качество DOCX (сложное форматирование упрощается). "
        "Для максимальной точности используйте pandoc."
    ]
    _render_blocks(doc, body, base_dir, Pt)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return BackendOutput(warnings=warnings)


_BLOCK_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "ul", "ol",
               "blockquote", "pre", "table", "hr", "img", "figure"}


def _render_blocks(doc, container, base_dir: Path, Pt, list_level: int = 0) -> None:
    from docx.shared import Inches

    for el in container.find_all(recursive=False):
        name = getattr(el, "name", None)
        if name is None:
            continue
        if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            doc.add_heading(el.get_text(strip=True), level=int(name[1]))
        elif name == "p":
            imgs = el.find_all("img")
            text = el.get_text(strip=True)
            if imgs and not text:
                for im in imgs:
                    _add_image(doc, im, base_dir, Inches)
            else:
                para = doc.add_paragraph()
                _add_inline_runs(para, el)
                for im in imgs:  # смешанный абзац: сохранить и текст, и картинки
                    _add_image(doc, im, base_dir, Inches)
        elif name in ("ul", "ol"):
            _render_list(doc, el, base_dir, Pt, ordered=(name == "ol"), level=list_level)
        elif name == "blockquote":
            para = doc.add_paragraph(style=_safe_style(doc, "Quote"))
            _add_inline_runs(para, el)
        elif name == "pre":
            text = el.get_text()
            para = doc.add_paragraph()
            run = para.add_run(text)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif name == "table":
            _render_table(doc, el)
        elif name == "img":
            _add_image(doc, el, base_dir, Inches)
        elif name == "figure":
            _render_blocks(doc, el, base_dir, Pt, list_level)
        elif name == "hr":
            doc.add_paragraph("―" * 20)
        else:
            text = el.get_text(strip=True)
            if text:
                doc.add_paragraph(text)


def _render_list(doc, el, base_dir, Pt, ordered: bool, level: int) -> None:
    base_style = "List Number" if ordered else "List Bullet"
    style_name = base_style if level == 0 else f"{base_style} {min(level + 1, 3)}"
    for li in el.find_all("li", recursive=False):
        para = doc.add_paragraph(style=_safe_style(doc, style_name))
        _add_inline_runs(para, li, skip_nested_lists=True)
        for sub in li.find_all(["ul", "ol"], recursive=False):
            _render_list(doc, sub, base_dir, Pt, ordered=(sub.name == "ol"), level=level + 1)


def _render_table(doc, el) -> None:
    rows = el.find_all("tr")
    if not rows:
        return
    ncols = max(len(r.find_all(["td", "th"])) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = _safe_table_style(doc, "Table Grid")
    for r in rows:
        cells = r.find_all(["td", "th"])
        row_cells = table.add_row().cells
        for i, cell in enumerate(cells):
            if i < ncols:
                row_cells[i].text = cell.get_text(strip=True)


def _add_inline_runs(para, element, skip_nested_lists: bool = False) -> None:
    for text, bold, italic, code in _iter_runs(element, skip_nested_lists=skip_nested_lists):
        if not text:
            continue
        run = para.add_run(text)
        run.bold = bold or None
        run.italic = italic or None
        if code:
            run.font.name = "Consolas"


def _iter_runs(node, bold=False, italic=False, code=False, skip_nested_lists=False):
    from bs4 import NavigableString

    for child in node.children:
        if isinstance(child, NavigableString):
            yield (str(child), bold, italic, code)
            continue
        name = getattr(child, "name", None)
        if name is None:
            continue
        if skip_nested_lists and name in ("ul", "ol"):
            continue
        b = bold or name in ("strong", "b")
        i = italic or name in ("em", "i")
        c = code or name in ("code", "tt", "kbd", "samp")
        if name == "br":
            yield ("\n", bold, italic, code)
            continue
        yield from _iter_runs(child, b, i, c, skip_nested_lists)


def _add_image(doc, img, base_dir: Path, Inches) -> None:
    src = img.get("src")
    if not src or src.startswith(("http://", "https://", "data:")):
        return
    path = (base_dir / src).resolve()
    if not path.exists():
        return
    try:
        doc.add_picture(str(path), width=Inches(5.5))
    except Exception:  # noqa: BLE001 — не роняем конвертацию из-за одной картинки
        pass


def _safe_style(doc, name: str) -> str:
    try:
        _ = doc.styles[name]
        return name
    except KeyError:
        return "Normal"


def _safe_table_style(doc, name: str):
    try:
        _ = doc.styles[name]
        return name
    except KeyError:
        return None


# --------------------------------------------------------------------------- #
# * -> PDF  (через HTML + xhtml2pdf)
# --------------------------------------------------------------------------- #
_PDF_SOURCES = {Format.MARKDOWN, Format.HTML, Format.DOCX, Format.EPUB}


def _to_pdf(inp: Path, out_path: Path, src: Format, o: ConvertOptions) -> BackendOutput:
    warnings: List[str] = []
    if src == Format.MARKDOWN:
        body = _md_to_html_str(_read_text(inp))
        base_dir = inp.parent
    elif src == Format.HTML:
        body = str(_html_body(_read_text(inp)))
        base_dir = inp.parent
    elif src == Format.DOCX:
        import mammoth

        with open(inp, "rb") as fh:
            result = mammoth.convert_to_html(fh)  # картинки -> data-URI по умолчанию
        body = result.value
        base_dir = inp.parent
    elif src == Format.EPUB:
        body = _epub_to_html_inline(inp)
        base_dir = inp.parent
    else:  # pragma: no cover
        raise ConversionFailedError(f"PDF из {src.title} не поддерживается.")

    user_css = _read_text(Path(o.css)) if o.css and Path(o.css).exists() else None
    title = _first_heading(body) or inp.stem
    warnings += render_html_to_pdf(
        body, out_path, base_dir=base_dir, title=title,
        page_size=o.pdf_page_size, user_css=user_css,
    )
    return BackendOutput(warnings=warnings)


def _epub_to_html_inline(inp: Path) -> str:
    """EPUB -> одна HTML-строка с картинками, встроенными как data-URI."""
    import base64

    ebooklib, epub, book = _load_epub(inp)
    images: Dict[str, str] = {}
    for item in book.get_items_of_type(ebooklib.ITEM_IMAGE):
        full = _norm_epub_path(item.get_name())
        b64 = base64.b64encode(item.get_content()).decode("ascii")
        mtype = _guess_media_type(Path(full).suffix)
        images[full] = f"data:{mtype};base64,{b64}"

    parts: List[str] = []
    for item in _spine_documents(ebooklib, book):
        body = _html_body(item.get_content().decode("utf-8", "replace"))
        for img in body.find_all("img"):
            src = img.get("src")
            if not src or src.startswith(("http://", "https://", "data:")):
                continue
            uri = _lookup_by_path(images, _resolve_doc_img(item.get_name(), src))
            if uri:
                img["src"] = uri
        parts.append(str(body))
    return "\n".join(parts)


# --------------------------------------------------------------------------- #
# Утилиты
# --------------------------------------------------------------------------- #
def _first_heading(html: str) -> Optional[str]:
    soup = BeautifulSoup(html, "html.parser")
    h = soup.find(["h1", "h2", "title"])
    return h.get_text(strip=True) if h else None


def _guess_media_type(suffix: str) -> str:
    suffix = suffix.lower()
    for mtype, ext in _IMG_EXT.items():
        if ext == suffix:
            return mtype
    return "image/png"


def _clean_md(text: str) -> str:
    # Схлопнуть 3+ пустых строки до одной пустой.
    lines = text.replace("\r\n", "\n").split("\n")
    out: List[str] = []
    blanks = 0
    for line in lines:
        if line.strip() == "":
            blanks += 1
            if blanks <= 1:
                out.append("")
        else:
            blanks = 0
            out.append(line.rstrip())
    return "\n".join(out).strip() + "\n"


# Диспетчер: (источник, цель) -> функция. Определяется после функций.
_DISPATCH = {
    (Format.MARKDOWN, Format.HTML): _md_to_html,
    (Format.HTML, Format.MARKDOWN): _html_to_md,
    (Format.DOCX, Format.MARKDOWN): _docx_to_md,
    (Format.EPUB, Format.MARKDOWN): _epub_to_md,
    (Format.MARKDOWN, Format.EPUB): _md_to_epub,
    (Format.MARKDOWN, Format.DOCX): _md_to_docx,
    (Format.HTML, Format.DOCX): _html_to_docx,
}
